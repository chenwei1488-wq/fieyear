import os
import sqlite3
try:
    import libsql_experimental as libsql
except ImportError:
    libsql = None
import json
import time
import datetime
import tempfile
import urllib.request
import urllib.parse
import hashlib
import io
import shutil
import openpyxl
from github import Github

from urllib.parse import urlparse, parse_qs
from jinja2 import Environment, FileSystemLoader

import boto3
from botocore.config import Config
from PIL import Image
# from rembg import remove
from dotenv import load_dotenv

# 加载 .env 文件
load_dotenv()

# ==================== 全局配置中心 ====================
DB_NAME = 'fiyear_data.db'
RESOURCE_DIR = 'resources'

R2_ENDPOINT = os.getenv('R2_ENDPOINT')
R2_ACCESS_KEY = os.getenv('R2_ACCESS_KEY')
R2_SECRET_KEY = os.getenv('R2_SECRET_KEY')
R2_BUCKET = 'fieyear-images'
CDN_DOMAIN = 'https://pub-bab8f5485b0242378eabd7ee02fcf2b0.r2.dev'

# ==================== 1. 纯净图像处理工具箱 ====================
class ImageToolbox:
    @staticmethod
    def process_remove_bg(input_path):
        if input_path.lower().startswith(('http://', 'https://')):
            req = urllib.request.Request(input_path, headers={'User-Agent': 'Mozilla/5.0'})
            with urllib.request.urlopen(req, timeout=15) as response:
                input_data = response.read()
        else:
            with open(input_path, 'rb') as i:
                input_data = i.read()

        output_data = remove(input_data)
        img = Image.open(io.BytesIO(output_data))

        temp_dir = os.path.join(RESOURCE_DIR, 'temp_nobg')
        os.makedirs(temp_dir, exist_ok=True)
        
        safe_name = f"nobg_{hashlib.md5(input_path.encode()).hexdigest()[:10]}.png"
        output_path = os.path.join(temp_dir, safe_name)
        
        img.save(output_path, format="PNG")
        return output_path

    @staticmethod
    def download_thumbnail(url):
        if not url or not url.startswith('http'):
            return None
            
        thumb_dir = os.path.join(RESOURCE_DIR, 'thumbnails')
        os.makedirs(thumb_dir, exist_ok=True)
        
        ext = url.split('.')[-1].split('?')[0]
        if len(ext) > 4: ext = 'jpg'
        safe_name = f"{hashlib.md5(url.encode()).hexdigest()}_thumb.{ext}"
        local_path = os.path.join(thumb_dir, safe_name)
        
        try:
            if not os.path.exists(local_path):
                req = urllib.request.Request(url, headers={'User-Agent': 'Mozilla/5.0'})
                with urllib.request.urlopen(req, timeout=5) as response, open(local_path, 'wb') as f:
                    f.write(response.read())
            return local_path
        except Exception:
            if os.path.exists(local_path):
                os.remove(local_path)
            return None

# ==================== 2. 核心数据与业务逻辑引擎 ====================
class BackendEngine:
    def __init__(self):
        self.init_db()

    def init_db(self):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        cursor.execute('''CREATE TABLE IF NOT EXISTS influencers (
                            id INTEGER PRIMARY KEY AUTOINCREMENT,
                            name TEXT, social_link TEXT, web_name TEXT, web_link TEXT,
                            platform TEXT DEFAULT 'Lovegobuy',
                            invite_code TEXT, invite_link TEXT, tutorial_link TEXT,
                            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''')
                            
        try: cursor.execute("ALTER TABLE influencers ADD COLUMN platform TEXT DEFAULT 'Lovegobuy'")
        except Exception: pass
        try: cursor.execute("ALTER TABLE influencers ADD COLUMN invite_btn_text TEXT")
        except Exception: pass
        try: cursor.execute("ALTER TABLE influencers ADD COLUMN tutorial_btn_text TEXT")
        except Exception: pass
        
        try: cursor.execute("ALTER TABLE influencers ADD COLUMN role TEXT DEFAULT 'tier2'")
        except Exception: pass
        try: cursor.execute("ALTER TABLE influencers ADD COLUMN parent_id INTEGER DEFAULT 0")
        except Exception: pass
                            
        cursor.execute('''CREATE TABLE IF NOT EXISTS products (
                            id INTEGER PRIMARY KEY AUTOINCREMENT,
                            influencer_id INTEGER, category TEXT, name TEXT,
                            product_link TEXT, transformed_link TEXT,
                            rmb_price TEXT, usd_price TEXT, eur_price TEXT, image_url TEXT,
                            qc_urls TEXT)''')
                            
        try: cursor.execute("ALTER TABLE products ADD COLUMN qc_urls TEXT")
        except Exception: pass
        try: cursor.execute("ALTER TABLE products ADD COLUMN sort_order INTEGER DEFAULT 0")
        except Exception: pass
        try: cursor.execute("ALTER TABLE products ADD COLUMN source_link TEXT")
        except Exception: pass
        try: cursor.execute("ALTER TABLE products ADD COLUMN sub_category TEXT DEFAULT ''")
        except Exception: pass
        try: cursor.execute("ALTER TABLE products ADD COLUMN created_at TEXT")
        except Exception: pass
        try: cursor.execute("ALTER TABLE products ADD COLUMN parent_product_id INTEGER DEFAULT 0")
        except Exception: pass
        try: cursor.execute("ALTER TABLE products ADD COLUMN status TEXT DEFAULT 'active'")
        except Exception: pass
        
        # 【新增】：tiktok视频ID，使用逗号分隔存储
        try: cursor.execute("ALTER TABLE products ADD COLUMN tiktok_video_ids TEXT DEFAULT ''")
        except Exception: pass

        cursor.execute("SELECT count(*) FROM sqlite_master WHERE type='table' AND name='categories_v2'")
        if cursor.fetchone()[0] == 0:
            cursor.execute('''CREATE TABLE categories_v2 (
                                influencer_id INTEGER,
                                name TEXT,
                                sort_order INTEGER DEFAULT 0,
                                UNIQUE(influencer_id, name)
                              )''')
            
            cursor.execute("SELECT count(*) FROM sqlite_master WHERE type='table' AND name='categories'")
            if cursor.fetchone()[0] > 0:
                cursor.execute("SELECT id FROM influencers")
                inf_ids = [r[0] for r in cursor.fetchall()]
                cursor.execute("SELECT name, sort_order FROM categories WHERE parent_name='' OR parent_name IS NULL")
                old_cats = cursor.fetchall()
                for i_id in inf_ids:
                    for c_name, c_sort in old_cats:
                        cursor.execute("INSERT OR IGNORE INTO categories_v2 (influencer_id, name, sort_order) VALUES (?, ?, ?)", (i_id, c_name, c_sort))

        cursor.execute('''CREATE TABLE IF NOT EXISTS sub_categories (
                            influencer_id INTEGER,
                            parent_category TEXT,
                            name TEXT,
                            sort_order INTEGER DEFAULT 0,
                            UNIQUE(influencer_id, parent_category, name)
                          )''')

        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        cursor.execute("UPDATE products SET created_at = ? WHERE created_at IS NULL OR created_at = ''", (now_str,))

        # ====== 新增：二手交易 Resell 数据表 ======
        cursor.execute('''CREATE TABLE IF NOT EXISTS resell_posts (
                            id INTEGER PRIMARY KEY AUTOINCREMENT,
                            name TEXT,
                            description TEXT,
                            image_url TEXT,
                            price TEXT,
                            purchase_link TEXT,
                            seller_discord TEXT,
                            status TEXT DEFAULT 'pending',
                            created_at TIMESTAMP DEFAULT CURRENT_TIMESTAMP)''')
        
        conn.commit()
        conn.close()

    # ==================== 二手审核专属接口 ====================
    def get_resell_posts(self, status=None):
        # ====== 新增：每次刷新时，从 Cloudflare Worker 拉取最新数据 ======
        import urllib.request
        import json
        
        # 你的专属 Worker 链接
        WORKER_URL = "https://resell-api.ww-4b9.workers.dev" 
        ADMIN_KEY = "admin123"
        
        try:
            # 1. 拉取请求：加入 User-Agent 伪装成浏览器，防止 403 拦截
            req = urllib.request.Request(
                f"{WORKER_URL}/pull?admin_key={ADMIN_KEY}",
                headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'}
            )
            with urllib.request.urlopen(req, timeout=5) as response:
                remote_posts = json.loads(response.read().decode())
                
            if remote_posts:
                conn = self.get_db_connection()
                cursor = conn.cursor()
                for p in remote_posts:
                    # 插入本地数据库
                    cursor.execute('''INSERT INTO resell_posts 
                                      (name, description, image_url, price, purchase_link, seller_discord, status)
                                      VALUES (?, ?, ?, ?, ?, ?, 'pending')''',
                                   (p.get('name'), p.get('desc'), p.get('img'), p.get('price'), p.get('link'), p.get('discord')))
                    
                    # 2. 删除请求：插入成功后告诉云端删除记录，同样需要加 headers
                    del_req = urllib.request.Request(
                        f"{WORKER_URL}/delete?admin_key={ADMIN_KEY}", 
                        data=json.dumps({"kv_key": p['kv_key']}).encode(),
                        headers={
                            'Content-Type': 'application/json',
                            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36'
                        }
                    )
                    urllib.request.urlopen(del_req, timeout=5)
                conn.commit()
                conn.close()
        except Exception as e:
            print("从云端同步二手帖子失败:", e)
            
        # ==========================================================
        
        # 原有的读取本地数据库逻辑
        conn = self.get_db_connection()
        cursor = conn.cursor()
        if status:
            cursor.execute("SELECT * FROM resell_posts WHERE status=? ORDER BY id DESC", (status,))
        else:
            cursor.execute("SELECT * FROM resell_posts ORDER BY CASE status WHEN 'pending' THEN 1 WHEN 'approved' THEN 2 ELSE 3 END, id DESC")
        res = cursor.fetchall()
        conn.close()
        return res

    def update_resell_status(self, post_id, new_status):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        cursor.execute("UPDATE resell_posts SET status=? WHERE id=?", (new_status, post_id))
        conn.commit()
        conn.close()
        return True

    def sync_tier1_to_new_tier2(self, tier2_inf_id, tier1_inf_id):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        
        cursor.execute("SELECT id FROM products WHERE influencer_id=? AND status='active'", (tier1_inf_id,))
        t1_products = cursor.fetchall()
        
        cursor.execute("SELECT parent_product_id FROM products WHERE influencer_id=?", (tier2_inf_id,))
        existing_parent_ids = {row[0] for row in cursor.fetchall() if row[0]}
        
        conn.close()
        
        for (t1_pid,) in t1_products:
            if t1_pid not in existing_parent_ids:
                self.push_product_to_tier2(t1_pid, tier2_inf_id)

    def push_product_to_tier2(self, tier1_pid, tier2_inf_id):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        
        cursor.execute("SELECT category, name, rmb_price, usd_price, eur_price, source_link, sub_category FROM products WHERE id=?", (tier1_pid,))
        tier1_product = cursor.fetchone()
        if not tier1_product:
            conn.close()
            return False
            
        c_cat, c_name, c_rmb, c_usd, c_eur, c_source, c_subcat = tier1_product
        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        
        # 注意：此处不插入 tiktok_video_ids，默认保持空白，从而隔离公库和私库的视频ID
        cursor.execute('''INSERT INTO products
                          (influencer_id, category, name, product_link, transformed_link,
                           rmb_price, usd_price, eur_price, image_url, qc_urls,
                           source_link, sub_category, created_at, parent_product_id, status)
                          VALUES (?, ?, ?, '', '', ?, ?, ?, '', '', ?, ?, ?, ?, 'pending')''',
                       (tier2_inf_id, c_cat, c_name, c_rmb, c_usd, c_eur, c_source, c_subcat, now_str, tier1_pid))
                       
        if c_cat:
            cursor.execute("INSERT OR IGNORE INTO categories_v2 (influencer_id, name) VALUES (?, ?)", (tier2_inf_id, c_cat))
            if c_subcat:
                cursor.execute("INSERT OR IGNORE INTO sub_categories (influencer_id, parent_category, name) VALUES (?, ?, ?)", (tier2_inf_id, c_cat, c_subcat))
            
        conn.commit()
        conn.close()

        return True

    def get_pending_tasks(self, tier2_inf_id):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        cursor.execute("""SELECT p.id, p.name, p.category, p.rmb_price, p.parent_product_id, parent.product_link
                          FROM products p
                          LEFT JOIN products parent ON p.parent_product_id = parent.id
                          WHERE p.influencer_id=? AND p.status='pending'
                          ORDER BY p.id ASC""", (tier2_inf_id,))
        res = cursor.fetchall()
        conn.close()
        return res

    def resolve_pending_task(self, tier2_pid, weidian_link):
        if not weidian_link: return False
        
        conn = self.get_db_connection()
        cursor = conn.cursor()
        
        cursor.execute("SELECT influencer_id, parent_product_id FROM products WHERE id=? AND status='pending'", (tier2_pid,))
        task_info = cursor.fetchone()
        if not task_info:
            conn.close()
            return False
            
        tier2_inf_id, parent_pid = task_info
        
        cursor.execute("SELECT image_url, qc_urls, source_link FROM products WHERE id=?", (parent_pid,))
        parent_media = cursor.fetchone()
        parent_img_url = parent_media[0] if parent_media else ""
        parent_qc_urls = parent_media[1] if parent_media else ""
        parent_source_link = parent_media[2] if parent_media else ""
        
        cursor.execute("SELECT invite_code, platform FROM influencers WHERE id=?", (tier2_inf_id,))
        inf_row = cursor.fetchone()
        inv_code = inf_row[0] if inf_row and inf_row[0] else ""
        plat = inf_row[1] if inf_row and inf_row[1] else "Lovegobuy"
        transformed_link = self.convert_link(weidian_link, plat, inv_code)
        
        cursor.execute('''UPDATE products
                          SET product_link=?, transformed_link=?, image_url=?, qc_urls=?, source_link=?, status='active'
                          WHERE id=?''',
                       (weidian_link, transformed_link, parent_img_url, parent_qc_urls, parent_source_link, tier2_pid))
                       
        conn.commit()
        conn.close()
        return True

    def get_combined_products_for_export(self, inf_id):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        
        # 加入了 p.tiktok_video_ids
        query = """
            SELECT p.id, p.influencer_id, p.category, p.name, p.transformed_link,
                   p.rmb_price, p.usd_price, p.eur_price, p.image_url, p.qc_urls, p.source_link, p.sub_category, p.created_at, p.tiktok_video_ids
            FROM products p
            LEFT JOIN categories_v2 c ON p.category = c.name AND p.influencer_id = c.influencer_id
            WHERE p.influencer_id = ? AND p.status = 'active'
            ORDER BY c.sort_order ASC, p.category ASC, p.sort_order ASC, p.id DESC
        """
        cursor.execute(query, (inf_id,))
        res = cursor.fetchall()
        conn.close()
        return res

    def get_categories(self, inf_id):
        if not inf_id: return []
        conn = self.get_db_connection()
        c = conn.cursor()
        c.execute("SELECT name FROM categories_v2 WHERE influencer_id=? ORDER BY sort_order ASC, name ASC", (inf_id,))
        rows = c.fetchall()
        conn.close()
        return [{"name": row[0]} for row in rows]

    def get_categories_tree(self, inf_id):
        if not inf_id: return []
        conn = self.get_db_connection()
        c = conn.cursor()
        c.execute("SELECT name FROM categories_v2 WHERE influencer_id=? ORDER BY sort_order ASC, name ASC", (inf_id,))
        main_cats = c.fetchall()
        
        tree = []
        for (m_cat,) in main_cats:
            c.execute("SELECT name FROM sub_categories WHERE influencer_id=? AND parent_category=? ORDER BY sort_order ASC, name ASC", (inf_id, m_cat))
            sub_cats = [{"name": row[0]} for row in c.fetchall()]
            tree.append({
                "name": m_cat,
                "subCategories": sub_cats
            })
            
        conn.close()
        return tree

    def update_categories_order(self, inf_id, ordered_names):
        if not ordered_names or not inf_id: return
        conn = self.get_db_connection()
        c = conn.cursor()
        for idx, name in enumerate(ordered_names):
            c.execute("UPDATE categories_v2 SET sort_order=? WHERE influencer_id=? AND name=?", (idx, inf_id, name))
        conn.commit()
        conn.close()

    def add_category(self, inf_id, name):
        if not name or not inf_id: return
        conn = self.get_db_connection()
        c = conn.cursor()
        c.execute("INSERT OR IGNORE INTO categories_v2 (influencer_id, name) VALUES (?, ?)", (inf_id, name))
        conn.commit()
        conn.close()
        
    def rename_category(self, inf_id, old_name, new_name):
        if not new_name or old_name == new_name or not inf_id: return
        conn = self.get_db_connection()
        c = conn.cursor()
        try:
            c.execute("UPDATE categories_v2 SET name=? WHERE influencer_id=? AND name=?", (new_name, inf_id, old_name))
            c.execute("UPDATE products SET category=? WHERE influencer_id=? AND category=?", (new_name, inf_id, old_name))
            c.execute("UPDATE sub_categories SET parent_category=? WHERE influencer_id=? AND parent_category=?", (new_name, inf_id, old_name))
            conn.commit()
        except sqlite3.IntegrityError:
            pass
        conn.close()

    def remove_category(self, inf_id, name):
        if not inf_id: return
        conn = self.get_db_connection()
        c = conn.cursor()
        c.execute("DELETE FROM categories_v2 WHERE influencer_id=? AND name=?", (inf_id, name))
        c.execute("UPDATE products SET category='' WHERE influencer_id=? AND category=?", (inf_id, name))
        c.execute("DELETE FROM sub_categories WHERE influencer_id=? AND parent_category=?", (inf_id, name))
        c.execute("UPDATE products SET sub_category='' WHERE influencer_id=? AND category=?", (inf_id, name))
        conn.commit()
        conn.close()

    # ==================== 二级类目专属操作接口 ====================
    def get_sub_categories(self, inf_id, parent_category):
        if not inf_id or not parent_category: return []
        conn = self.get_db_connection()
        c = conn.cursor()
        c.execute("SELECT name FROM sub_categories WHERE influencer_id=? AND parent_category=? ORDER BY sort_order ASC, name ASC", (inf_id, parent_category))
        rows = c.fetchall()
        conn.close()
        return [{"name": row[0]} for row in rows]

    def add_sub_category(self, inf_id, parent_category, name):
        if not name or not inf_id or not parent_category: return
        conn = self.get_db_connection()
        c = conn.cursor()
        c.execute("INSERT OR IGNORE INTO sub_categories (influencer_id, parent_category, name) VALUES (?, ?, ?)", (inf_id, parent_category, name))
        conn.commit()
        conn.close()

    def rename_sub_category(self, inf_id, parent_category, old_name, new_name):
        if not new_name or old_name == new_name or not inf_id or not parent_category: return
        conn = self.get_db_connection()
        c = conn.cursor()
        try:
            c.execute("UPDATE sub_categories SET name=? WHERE influencer_id=? AND parent_category=? AND name=?", (new_name, inf_id, parent_category, old_name))
            c.execute("UPDATE products SET sub_category=? WHERE influencer_id=? AND category=? AND sub_category=?", (new_name, inf_id, parent_category, old_name))
            conn.commit()
        except sqlite3.IntegrityError:
            pass
        conn.close()

    def remove_sub_category(self, inf_id, parent_category, name):
        if not inf_id or not parent_category: return
        conn = self.get_db_connection()
        c = conn.cursor()
        c.execute("DELETE FROM sub_categories WHERE influencer_id=? AND parent_category=? AND name=?", (inf_id, parent_category, name))
        c.execute("UPDATE products SET sub_category='' WHERE influencer_id=? AND category=? AND sub_category=?", (inf_id, parent_category, name))
        conn.commit()
        conn.close()

    def update_sub_categories_order(self, inf_id, parent_category, ordered_names):
        if not ordered_names or not inf_id or not parent_category: return
        conn = self.get_db_connection()
        c = conn.cursor()
        for idx, name in enumerate(ordered_names):
            c.execute("UPDATE sub_categories SET sort_order=? WHERE influencer_id=? AND parent_category=? AND name=?", (idx, inf_id, parent_category, name))
        conn.commit()
        conn.close()
    # ==========================================================

    def clean_price_to_string(self, text):
        if text is None: return "0.00"
        clean_str = str(text).replace(',', '.')
        clean_str = clean_str.replace('€', '').replace('$', '').replace('¥', '').replace('￥', '').strip()
        if clean_str.count('.') > 1:
            parts = clean_str.split('.')
            clean_str = "".join(parts[:-1]) + "." + parts[-1]
        try:
            val = float(clean_str)
            return f"{val:.2f}"
        except ValueError:
            return "0.00"

    def extract_item_id(self, url):
        try:
            parsed = urlparse(url)
            params = parse_qs(parsed.query)
            item_id = params.get('itemID', [''])[0] or params.get('id', [''])[0]
            if item_id: return item_id
            return hashlib.md5(url.encode('utf-8')).hexdigest()[:12]
        except:
            return str(int(time.time()))

    def add_influencer(self, data):
        if len(data) == 10:
            data = (*data, 'tier2', 0)
            
        conn = self.get_db_connection()
        cursor = conn.cursor()
        cursor.execute('''INSERT INTO influencers
                          (name, social_link, web_name, web_link, platform, invite_code, invite_link, invite_btn_text, tutorial_link, tutorial_btn_text, role, parent_id)
                          VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''', data)
        new_id = cursor.lastrowid
        conn.commit()
        conn.close()
        
        if len(data) >= 12 and data[10] == 'tier2' and int(data[11]) > 0:
            self.sync_tier1_to_new_tier2(new_id, int(data[11]))
            
        return new_id

    def update_influencer(self, inf_id, data):
        if len(data) == 10:
            data = (*data, 'tier2', 0)
            
        conn = self.get_db_connection()
        cursor = conn.cursor()
        
        cursor.execute('''UPDATE influencers SET
                          name=?, social_link=?, web_name=?, web_link=?, platform=?,
                          invite_code=?, invite_link=?, invite_btn_text=?, tutorial_link=?, tutorial_btn_text=?, role=?, parent_id=? WHERE id=?''', (*data, inf_id))
        
        platform = data[4]
        invite_code = data[5]
        
        cursor.execute("SELECT id, product_link FROM products WHERE influencer_id=?", (inf_id,))
        products = cursor.fetchall()
        for pid, product_link in products:
            if product_link:
                new_transformed_link = self.convert_link(product_link, platform, invite_code)
                cursor.execute("UPDATE products SET transformed_link=? WHERE id=?", (new_transformed_link, pid))
                
        conn.commit()
        conn.close()

        if len(data) >= 12 and data[10] == 'tier2' and int(data[11]) > 0:
            self.sync_tier1_to_new_tier2(inf_id, int(data[11]))

    def delete_influencer(self, inf_id):
        products = self.get_products_by_influencers([inf_id])
        if products:
            product_ids = [p[0] for p in products]
            self.delete_products(product_ids)
            
        conn = self.get_db_connection()
        cursor = conn.cursor()
        cursor.execute("DELETE FROM influencers WHERE id = ?", (inf_id,))
        cursor.execute("DELETE FROM categories_v2 WHERE influencer_id = ?", (inf_id,))
        cursor.execute("DELETE FROM sub_categories WHERE influencer_id = ?", (inf_id,))
        conn.commit()
        conn.close()

    def import_products_from_excel(self, excel_path, influencer_id, auto_remove_bg=False, progress_cb=None):
        if not os.path.exists(excel_path) or not excel_path.endswith(('.xlsx', '.xls')):
            return 0
        try:
            wb = openpyxl.load_workbook(excel_path, data_only=True)
            sheet = wb.active
        except Exception:
            return 0

        conn = self.get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT invite_code, platform FROM influencers WHERE id=?", (influencer_id,))
        inf_row = cursor.fetchone()
        conn.close()
        invite_code = inf_row[0] if inf_row and inf_row[0] else ""
        platform = inf_row[1] if inf_row and inf_row[1] else "Lovegobuy"

        rows = list(sheet.iter_rows(min_row=2, values_only=True))
        total_rows = len(rows)
        success_count = 0

        for idx, row in enumerate(rows):
            name = str(row[0]).strip() if len(row) > 0 and row[0] is not None else "未命名商品"
            image_url = str(row[1]).strip() if len(row) > 1 and row[1] is not None else ""
            category = str(row[2]).strip() if len(row) > 2 and row[2] is not None else ""
            product_link = str(row[3]).strip() if len(row) > 3 and row[3] is not None else ""
            rmb_price = str(row[4]).strip() if len(row) > 4 and row[4] is not None else "0.00"
            
            qc_urls = str(row[5]).strip() if len(row) > 5 and row[5] is not None else ""
            source_link = str(row[6]).strip() if len(row) > 6 and row[6] is not None else ""

            if not product_link or not product_link.startswith('http'):
                if progress_cb: progress_cb(idx + 1, total_rows)
                continue
                
            transformed_link = self.convert_link(product_link, platform, invite_code)
            
            is_r2_img = "pub-bab8f5485b0242378eabd7ee02fcf2b0.r2.dev" in image_url
            initial_img = image_url if is_r2_img else ""
                
            # 加入默认 tk_id 的空字符串
            product_data = [
                influencer_id, category, name, product_link, transformed_link,
                rmb_price, "0.00", "0.00", initial_img, qc_urls, source_link, "", ""
            ]
            
            try:
                new_id = self.add_single_product_get_id(product_data)
                item_id = self.extract_item_id(product_link)
                
                if category:
                    self.add_category(influencer_id, category)
                
                if image_url and not is_r2_img:
                    cdn_url = self.process_image_pipeline(
                        image_url, influencer_id, item_id,
                        is_qc=False, remove_bg=auto_remove_bg
                    )
                    if cdn_url:
                        self.update_product_image(new_id, cdn_url)
                        
                success_count += 1
            except Exception:
                pass
            
            if progress_cb: progress_cb(idx + 1, total_rows)

        return success_count

    def add_single_product_get_id(self, data):
        cleaned_data = list(data)
        if len(cleaned_data) == 11: cleaned_data.append("")  # sub_category
        if len(cleaned_data) == 12: cleaned_data.append("")  # tiktok_video_ids
            
        cleaned_data[5] = self.clean_price_to_string(cleaned_data[5])
        cleaned_data[6] = self.clean_price_to_string(cleaned_data[6])
        cleaned_data[7] = self.clean_price_to_string(cleaned_data[7])
        
        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        cleaned_data.append(now_str)

        conn = self.get_db_connection()
        cursor = conn.cursor()
        cursor.execute('''INSERT INTO products
                          (influencer_id, category, name, product_link, transformed_link, rmb_price, usd_price, eur_price, image_url, qc_urls, source_link, sub_category, tiktok_video_ids, created_at)
                          VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)''', cleaned_data)
        new_id = cursor.lastrowid
        conn.commit()

        inf_id = cleaned_data[0]
        cursor.execute("SELECT role FROM influencers WHERE id=?", (inf_id,))
        role_row = cursor.fetchone()
        
        if role_row and role_row[0] == 'tier1':
            cursor.execute("SELECT id FROM influencers WHERE parent_id=? AND role='tier2'", (inf_id,))
            child_infs = cursor.fetchall()
            
            for child in child_infs:
                child_id = child[0]
                self.push_product_to_tier2(new_id, child_id)

        conn.close()
        return new_id

    def update_product_image(self, p_id, image_url):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        cursor.execute('UPDATE products SET image_url=? WHERE id=?', (image_url, p_id))
        conn.commit()
        conn.close()

    def update_product_qc_urls(self, p_id, qc_urls):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        cursor.execute('UPDATE products SET qc_urls=? WHERE id=?', (qc_urls, p_id))
        conn.commit()
        conn.close()

    def update_product(self, p_id, data):
        cleaned_data = list(data)
        if len(cleaned_data) == 11: cleaned_data.append("")
        if len(cleaned_data) == 12: cleaned_data.append("") # tiktok_video_ids
            
        cleaned_data[5] = self.clean_price_to_string(cleaned_data[5])
        cleaned_data[6] = self.clean_price_to_string(cleaned_data[6])
        cleaned_data[7] = self.clean_price_to_string(cleaned_data[7])

        conn = self.get_db_connection()
        cursor = conn.cursor()
        
        cursor.execute("SELECT status FROM products WHERE id=?", (p_id,))
        status_row = cursor.fetchone()
        current_status = status_row[0] if status_row else "active"
        
        new_status = 'active' if (current_status == 'pending' and cleaned_data[3]) else current_status

        cursor.execute('''UPDATE products
                          SET influencer_id=?, category=?, name=?, product_link=?, transformed_link=?, rmb_price=?, usd_price=?, eur_price=?, image_url=?, qc_urls=?, source_link=?, sub_category=?, tiktok_video_ids=?, status=?
                          WHERE id=?''', (*cleaned_data, new_status, p_id))
        
        # 级联更新私库，包含二级类目继承，但故意排除了 tiktok_video_ids （实现隔离）
        cursor.execute('''UPDATE products SET category=?, rmb_price=?, usd_price=?, eur_price=?, name=?, image_url=?, qc_urls=?, source_link=?, sub_category=? WHERE parent_product_id=?''',
                       (cleaned_data[1], cleaned_data[5], cleaned_data[6], cleaned_data[7], cleaned_data[2], cleaned_data[8], cleaned_data[9], cleaned_data[10], cleaned_data[11], p_id))
        
        if cleaned_data[1]:
            cursor.execute("SELECT DISTINCT influencer_id FROM products WHERE parent_product_id=?", (p_id,))
            child_infs = cursor.fetchall()
            for child in child_infs:
                cursor.execute("INSERT OR IGNORE INTO categories_v2 (influencer_id, name) VALUES (?, ?)", (child[0], cleaned_data[1]))
                if cleaned_data[11]:
                    cursor.execute("INSERT OR IGNORE INTO sub_categories (influencer_id, parent_category, name) VALUES (?, ?, ?)", (child[0], cleaned_data[1], cleaned_data[11]))
                
        conn.commit()
        conn.close()

    def recalculate_foreign_currencies(self, usd_rate, eur_rate, influencer_ids=None):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        
        if influencer_ids:
            placeholders = ','.join(['?'] * len(influencer_ids))
            query = f"SELECT id, rmb_price FROM products WHERE influencer_id IN ({placeholders})"
            cursor.execute(query, tuple(influencer_ids))
        else:
            cursor.execute("SELECT id, rmb_price FROM products")
            
        rows = cursor.fetchall()
        
        for pid, rmb in rows:
            try:
                rmb_f = float(rmb)
                new_usd = f"{(rmb_f / usd_rate):.2f}"
                new_eur = f"{(rmb_f / eur_rate):.2f}"
            except:
                new_usd, new_eur = "0.00", "0.00"
            cursor.execute("UPDATE products SET usd_price=?, eur_price=? WHERE id=?", (new_usd, new_eur, pid))
        conn.commit()
        conn.close()

    def update_sort_orders(self, ordered_ids):
        if not ordered_ids: return
        conn = self.get_db_connection()
        cursor = conn.cursor()
        for idx, pid in enumerate(ordered_ids):
            cursor.execute("UPDATE products SET sort_order=? WHERE id=?", (idx, pid))
        conn.commit()
        conn.close()

    def remove_product_from_category(self, p_id):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        cursor.execute("UPDATE products SET category='', sub_category='' WHERE id=?", (p_id,))
        conn.commit()
        conn.close()

    def change_product_category(self, p_id, new_category):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        cursor.execute("UPDATE products SET category=?, sub_category='' WHERE id=?", (new_category, p_id))
        conn.commit()
        conn.close()

    def bulk_update_products(self, p_ids, update_dict, usd_rate=7.14, eur_rate=7.69):
        conn = self.get_db_connection()
        cursor = conn.cursor()

        for pid in p_ids:
            # 加入了 tiktok_video_ids
            cursor.execute("SELECT influencer_id, category, name, product_link, rmb_price, usd_price, eur_price, image_url, transformed_link, source_link, sub_category, tiktok_video_ids FROM products WHERE id=?", (pid,))
            row = cursor.fetchone()
            if not row: continue
            c_inf, c_cat, c_name, c_plink, c_rmb, c_usd, c_eur, c_img, c_tlink, c_source, c_subcat, c_tk = row

            n_inf = update_dict.get('influencer_id', c_inf)
            n_cat = update_dict.get('category', c_cat)
            n_name = update_dict.get('name', c_name)
            n_source = update_dict.get('source_link', c_source)
            n_subcat = update_dict.get('sub_category', c_subcat)
            n_tk = update_dict.get('tiktok_video_ids', c_tk) # 新增获取
            
            if 'rmb_price' in update_dict:
                n_rmb = self.clean_price_to_string(update_dict['rmb_price'])
                try:
                    rmb_f = float(n_rmb)
                    n_usd = f"{(rmb_f / usd_rate):.2f}"
                    n_eur = f"{(rmb_f / eur_rate):.2f}"
                except:
                    n_usd, n_eur = "0.00", "0.00"
            else:
                n_rmb, n_usd, n_eur = c_rmb, c_usd, c_eur
            
            item_id = self.extract_item_id(c_plink)
            
            n_img = c_img
            if 'image_url' in update_dict:
                raw_img = update_dict['image_url']
                if raw_img: n_img = raw_img

            n_tlink = c_tlink
            if 'influencer_id' in update_dict:
                cursor.execute("SELECT invite_code, platform FROM influencers WHERE id=?", (n_inf,))
                inf_row = cursor.fetchone()
                inv_code = inf_row[0] if inf_row and inf_row[0] else ""
                plat = inf_row[1] if inf_row and inf_row[1] else "Lovegobuy"
                n_tlink = self.convert_link(c_plink, plat, inv_code)

            cursor.execute("SELECT qc_urls FROM products WHERE id=?", (pid,))
            c_qc_urls = cursor.fetchone()[0] or ""

            n_qc_urls = c_qc_urls
            if 'qc_urls' in update_dict:
                n_qc_urls = update_dict['qc_urls']

            # 更新语句加入 tiktok_video_ids=?
            cursor.execute('''UPDATE products SET influencer_id=?, category=?, name=?, transformed_link=?, rmb_price=?, usd_price=?, eur_price=?, image_url=?, qc_urls=?, source_link=?, sub_category=?, tiktok_video_ids=? WHERE id=?''',
                           (n_inf, n_cat, n_name, n_tlink, n_rmb, n_usd, n_eur, n_img, n_qc_urls, n_source, n_subcat, n_tk, pid))
            
            # 级联更新依然刻意排除 tiktok_video_ids，保证二开独立
            if 'rmb_price' in update_dict or 'name' in update_dict or 'image_url' in update_dict or 'qc_urls' in update_dict or 'source_link' in update_dict or 'category' in update_dict or 'sub_category' in update_dict:
                cursor.execute('''UPDATE products SET category=?, rmb_price=?, usd_price=?, eur_price=?, name=?, image_url=?, qc_urls=?, source_link=?, sub_category=? WHERE parent_product_id=?''',
                               (n_cat, n_rmb, n_usd, n_eur, n_name, n_img, n_qc_urls, n_source, n_subcat, pid))
                if n_cat:
                    cursor.execute("SELECT DISTINCT influencer_id FROM products WHERE parent_product_id=?", (pid,))
                    child_infs = cursor.fetchall()
                    for child in child_infs:
                        cursor.execute("INSERT OR IGNORE INTO categories_v2 (influencer_id, name) VALUES (?, ?)", (child[0], n_cat))
                        if n_subcat:
                            cursor.execute("INSERT OR IGNORE INTO sub_categories (influencer_id, parent_category, name) VALUES (?, ?, ?)", (child[0], n_cat, n_subcat))
            
            if n_cat:
                cursor.execute("INSERT OR IGNORE INTO categories_v2 (influencer_id, name) VALUES (?, ?)", (n_inf, n_cat))
            if n_cat and n_subcat:
                cursor.execute("INSERT OR IGNORE INTO sub_categories (influencer_id, parent_category, name) VALUES (?, ?, ?)", (n_inf, n_cat, n_subcat))

        conn.commit()
        conn.close()

    def delete_image_from_r2(self, remote_key):
        try:
            s3 = boto3.client('s3', endpoint_url=R2_ENDPOINT, aws_access_key_id=R2_ACCESS_KEY,
                              aws_secret_access_key=R2_SECRET_KEY, config=Config(signature_version='s3v4'))
            s3.delete_object(Bucket=R2_BUCKET, Key=remote_key)
        except Exception: pass

    def delete_products(self, product_ids):
        if not product_ids: return
        conn = self.get_db_connection()
        cursor = conn.cursor()
        
        placeholders = ','.join(['?'] * len(product_ids))
        cursor.execute(f"SELECT id FROM products WHERE parent_product_id IN ({placeholders})", tuple(product_ids))
        child_ids = [r[0] for r in cursor.fetchall()]
        all_ids_to_delete = list(product_ids) + child_ids
        
        if not all_ids_to_delete:
            return

        del_placeholders = ','.join(['?'] * len(all_ids_to_delete))
        cursor.execute(f"DELETE FROM products WHERE id IN ({del_placeholders})", tuple(all_ids_to_delete))
        conn.commit()
        conn.close()

    def get_influencers(self):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT id, name, social_link, web_name, web_link, platform, invite_code, invite_link, invite_btn_text, tutorial_link, tutorial_btn_text, role, parent_id FROM influencers")
        res = cursor.fetchall()
        conn.close()
        return res

    def get_product_by_id(self, p_id):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT * FROM products WHERE id=?", (p_id,))
        res = cursor.fetchone()
        conn.close()
        return res

    def get_products_by_influencers(self, inf_ids):
        if not inf_ids: return []
        conn = self.get_db_connection()
        cursor = conn.cursor()
        placeholders = ','.join(['?'] * len(inf_ids))
        
        # 加入了 p.tiktok_video_ids
        query = f"""
            SELECT p.id, p.influencer_id, p.category, p.name, p.transformed_link,
                   p.rmb_price, p.usd_price, p.eur_price, p.image_url, p.qc_urls, p.source_link, p.sub_category, p.created_at, p.tiktok_video_ids
            FROM products p
            LEFT JOIN categories_v2 c ON p.category = c.name AND p.influencer_id = c.influencer_id
            WHERE p.influencer_id IN ({placeholders}) AND p.status = 'active'
            ORDER BY c.sort_order ASC, p.category ASC, p.sort_order ASC, p.id DESC
        """
        cursor.execute(query, tuple(inf_ids))
        res = cursor.fetchall()
        conn.close()
        return res

    def search_products(self, inf_id, p_id="", name="", link="", start_date="", end_date="", cat=""):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        # 加入了 p.tiktok_video_ids
        query = """
            SELECT p.id, p.influencer_id, p.category, p.name, p.transformed_link,
                   p.rmb_price, p.usd_price, p.eur_price, p.image_url, p.qc_urls, p.source_link, p.sub_category, p.created_at, p.tiktok_video_ids
            FROM products p
            LEFT JOIN categories_v2 c ON p.category = c.name AND p.influencer_id = c.influencer_id
            WHERE p.influencer_id = ?
        """
        params = [inf_id]
        
        if p_id:
            query += " AND p.id = ?"
            params.append(p_id)
        if name:
            query += " AND p.name LIKE ?"
            params.append(f"%{name}%")
        if link:
            query += " AND (p.product_link LIKE ? OR p.transformed_link LIKE ? OR p.source_link LIKE ?)"
            params.extend([f"%{link}%", f"%{link}%", f"%{link}%"])
        if cat:
            query += " AND p.category = ?"
            params.append(cat)
        if start_date and end_date:
            query += " AND date(p.created_at) BETWEEN ? AND ?"
            params.extend([start_date, end_date])
            
        query += " ORDER BY p.id DESC"
        cursor.execute(query, tuple(params))
        res = cursor.fetchall()
        conn.close()
        return res
        
    def get_products_for_decoration(self, inf_id, category, sort_by="default"):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        order_clause = "ORDER BY sort_order ASC, id DESC"
        if sort_by == "alpha": order_clause = "ORDER BY name COLLATE NOCASE ASC"
        query = f"SELECT id, name, image_url, eur_price FROM products WHERE influencer_id=? AND category=? AND status='active' {order_clause}"
        cursor.execute(query, (inf_id, category))
        res = cursor.fetchall()
        conn.close()
        return res
   
    def update_product_order_by_ids(self, ordered_ids):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        for idx, pid in enumerate(ordered_ids):
            cursor.execute("UPDATE products SET sort_order=? WHERE id=?", (idx, pid))
        conn.commit()
        conn.close()

    def get_undecorated_products_by_influencer(self, inf_id):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT id, name, image_url, eur_price FROM products WHERE influencer_id=? AND (category IS NULL OR category='') AND status='active' ORDER BY id DESC", (inf_id,))
        res = cursor.fetchall()
        conn.close()
        return res

    def get_all_products_missing_qc(self):
        conn = self.get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT id, influencer_id, source_link, name FROM products WHERE (qc_urls IS NULL OR qc_urls = '') AND status='active'")
        res = cursor.fetchall()
        conn.close()
        return res

    def convert_link(self, original_url, platform, invite_code):
        try:
            if not original_url: return ""
            parsed = urlparse(original_url)
            params = parse_qs(parsed.query)
            item_id = params.get('itemID', [''])[0] or params.get('id', [''])[0]
            
            if item_id:
                if platform and platform.lower() == 'boonbuy':
                    code_str = ""
                    if invite_code:
                        invite_code = invite_code.strip()
                        if invite_code.startswith("?inviteCode="): code_str = invite_code
                        elif invite_code.startswith("?invite_code="): code_str = f"?inviteCode={invite_code[13:]}"
                        elif invite_code.startswith("&invite_code="): code_str = f"?inviteCode={invite_code[13:]}"
                        elif invite_code.startswith("inviteCode="): code_str = f"?{invite_code}"
                        else: code_str = f"?inviteCode={invite_code}"
                            
                    return f"https://www.boonbuy.com/product/weidian/{item_id}{code_str}"
                
                else:
                    code_str = ""
                    if invite_code:
                        invite_code = invite_code.strip()
                        if invite_code.startswith("&invite_code="): code_str = invite_code
                        elif invite_code.startswith("?inviteCode="): code_str = f"&invite_code={invite_code[12:]}"
                        elif invite_code.startswith("?invite_code="): code_str = f"&{invite_code[1:]}"
                        elif invite_code.startswith("invite_code="): code_str = f"&{invite_code}"
                        else: code_str = f"&invite_code={invite_code}"
                            
                    return f"https://www.lovegobuy.com/product?id={item_id}&shop_type=weidian{code_str}"
            return original_url
        except Exception:
            return original_url

    def process_image_pipeline(self, local_src_path, influencer_id, item_id, image_dir="", is_qc=False, qc_index=1, remove_bg=False):
        if not local_src_path: return None
        valid_path = None
        is_temp_file = False
        
        if local_src_path.lower().startswith(('http://', 'https://')):
            try:
                temp_fd, temp_path = tempfile.mkstemp(suffix='.jpg')
                os.close(temp_fd)
                req = urllib.request.Request(local_src_path, headers={'User-Agent': 'Mozilla/5.0'})
                with urllib.request.urlopen(req) as response, open(temp_path, 'wb') as out_file:
                    out_file.write(response.read())
                valid_path = temp_path
                is_temp_file = True
            except Exception:
                return None
        else:
            search_paths = [local_src_path]
            if image_dir:
                search_paths.append(os.path.join(image_dir, os.path.basename(local_src_path)))
            for path in search_paths:
                if os.path.exists(path):
                    valid_path = path
                    break
        if not valid_path: return None
        
        inf_dir = os.path.join(RESOURCE_DIR, str(influencer_id))
        if is_qc:
            sub_dir = os.path.join(inf_dir, "qc")
            os.makedirs(sub_dir, exist_ok=True)
            webp_filename = f"{item_id}_qc{qc_index}.webp"
            remote_key = f"{influencer_id}/qc/{webp_filename}"
            webp_local_path = os.path.join(sub_dir, webp_filename)
        else:
            sub_dir = os.path.join(inf_dir, "product")
            os.makedirs(sub_dir, exist_ok=True)
            webp_filename = f"{item_id}.webp"
            remote_key = f"{influencer_id}/product/{webp_filename}"
            webp_local_path = os.path.join(sub_dir, webp_filename)
            
        try:
            if remove_bg:
                with open(valid_path, 'rb') as i_file:
                    input_data = i_file.read()
                output_data = remove(input_data)
                img = Image.open(io.BytesIO(output_data)).convert("RGBA")
            else:
                img = Image.open(valid_path)
            
            img.save(webp_local_path, "WEBP", quality=80)
            
            s3 = boto3.client('s3', endpoint_url=R2_ENDPOINT, aws_access_key_id=R2_ACCESS_KEY,
                              aws_secret_access_key=R2_SECRET_KEY, config=Config(signature_version='s3v4'))
            s3.upload_file(webp_local_path, R2_BUCKET, remote_key)
            cdn_url = f"{CDN_DOMAIN.rstrip('/')}/{remote_key}"
            if is_temp_file and os.path.exists(valid_path): os.remove(valid_path)
            return cdn_url
        except Exception as e:
            print(f"Image Error: {e}")
            if is_temp_file and os.path.exists(valid_path): os.remove(valid_path)
            return None

    def get_db_connection(self):
        db_url = os.getenv('TURSO_DB_URL')
        auth_token = os.getenv('TURSO_AUTH_TOKEN')
        
        # 只有在有密钥，且云端库(libsql)存在时，才连接 Turso 云数据库
        if db_url and auth_token and libsql:
            conn = libsql.connect(database=db_url, auth_token=auth_token)
        else:
            # 在本地 Windows 开发时，自动退回使用本地的数据库文件
            conn = sqlite3.connect('fiyear_data.db')
        return conn

    def export_web_page(self, product_data_list, display_currency="EUR", export_dir=""):
        if not product_data_list: return False
        
        try:
            first_inf_id = product_data_list[0][8] if isinstance(product_data_list[0][8], int) else product_data_list[0][1]
        except IndexError:
            first_inf_id = product_data_list[0][1]
            
        conn = self.get_db_connection()
        cursor = conn.cursor()
        cursor.execute("SELECT name, invite_link, tutorial_link, invite_btn_text, tutorial_btn_text FROM influencers WHERE id=?", (first_inf_id,))
        inf = cursor.fetchone()
        conn.close()

        if not inf: return False
        
        influencer_name = inf[0] if inf[0] else "Fieyear"
        
        invite_link = inf[1] or "#"
        tutorial_link = inf[2] or "#"
        invite_btn_text = inf[3] if inf[3] else "🎁 S'inscrire sur Lovegobuy pour -30%"
        tutorial_btn_text = inf[4] if inf[4] else "📚 Tutoriel"

        os.makedirs(export_dir, exist_ok=True)
        timestamp = int(time.time())
        product_json_list = []
        symbol_map = {"EUR": "€", "USD": "$", "RMB": "¥"}
        symbol = symbol_map.get(display_currency, "€")
        
        sorted_categories_js = json.dumps(self.get_categories_tree(first_inf_id), ensure_ascii=False)

        for p in product_data_list:
            if display_currency == "RMB": price_val = p[5]
            elif display_currency == "USD": price_val = p[6]
            else: price_val = p[7]
            
            img_url_raw = str(p[8]) if len(p) > 8 and p[8] else ""
            clean_img_url = img_url_raw.split('?')[0] if img_url_raw else ""
            final_img_url = f"{clean_img_url}?v={timestamp}" if clean_img_url else ""
            
            qc_raw = str(p[9]) if len(p) > 9 and p[9] else ""
            qc_list = [f"{url.strip()}?v={timestamp}" for url in qc_raw.split(",") if url.strip()]
            
            if not p[2]: continue
            
            # 提取 tiktok 视频 ID，并支持多个逗号分隔转换成数组
            tk_raw = str(p[13]) if len(p) > 13 and p[13] else ""
            tk_list = [tk.strip() for tk in tk_raw.split(",") if tk.strip()]
            
            product_json_list.append({
                "id": p[0],
                "name": p[3],
                "productLink": p[4],
                "price": price_val,
                "symbol": symbol,
                "category": p[2],
                "subCategory": p[11] if len(p) > 11 and p[11] else "",
                "imgUrl": final_img_url,
                "qcUrls": qc_list,
                "tiktokVideoIds": tk_list  # 打包进最终页面的 JSON 数据源中
            })

        json_path = os.path.join(export_dir, "products.json")
        # ====== 新增：将已审核通过的 Resell 商品注入 JSON ======
        cursor = sqlite3.connect(DB_NAME).cursor()
        cursor.execute("SELECT id, name, purchase_link, price, image_url, seller_discord, description FROM resell_posts WHERE status='approved'")
        approved_resells = cursor.fetchall()
        for r in approved_resells:
            product_json_list.append({
                "id": f"resell_{r[0]}",
                "name": r[1],
                "productLink": r[2],  # 买家点击跳转的购买链接
                "price": r[3],
                "symbol": symbol,
                "category": "RESELL",
                "subCategory": "",
                "imgUrl": r[4],
                "sellerDiscord": r[5], # 卖家的联系方式
                "description": r[6],
                "container": "RESELL"  # 关键：告诉前端属于RESELL容器
            })
        with open(json_path, "w", encoding="utf-8") as f:
            f.write(json.dumps(product_json_list, ensure_ascii=False))

        # ====== 增强版：自动分发所有静态模板与资源 ======
        # 1. 拷贝 Logo
        if os.path.exists("logo.png"): shutil.copy("logo.png", os.path.join(export_dir, "logo.png"))
        
        # 2. 拷贝 CSS 样式表
        style_src = os.path.join("templates", "style.css")
        if os.path.exists(style_src): shutil.copy(style_src, os.path.join(export_dir, "style.css"))
        
        # 3. 自动拷贝 Cloudflare 的 _worker.js
        if os.path.exists("_worker.js"):
            shutil.copy("_worker.js", os.path.join(export_dir, "_worker.js"))
        elif os.path.exists(os.path.join("templates", "_worker.js")):
            shutil.copy(os.path.join("templates", "_worker.js"), os.path.join(export_dir, "_worker.js"))

        if os.path.exists("logo.png"): shutil.copy("logo.png", os.path.join(export_dir, "logo.png"))
        style_src = os.path.join("templates", "style.css")
        if os.path.exists(style_src): shutil.copy(style_src, os.path.join(export_dir, "style.css"))

        env = Environment(loader=FileSystemLoader('templates'), variable_start_string='[[', variable_end_string=']]')
        qc_template = env.get_template('qc.html')
        with open(os.path.join(export_dir, 'qc.html'), 'w', encoding='utf-8') as f: f.write(qc_template.render())

        index_template = env.get_template('index.html')
        with open(os.path.join(export_dir, 'index.html'), 'w', encoding='utf-8') as f:
            f.write(index_template.render(
                influencer_name=influencer_name,
                invite_link=invite_link,
                tutorial_link=tutorial_link,
                invite_btn_text=invite_btn_text,
                tutorial_btn_text=tutorial_btn_text,
                sorted_categories_js=sorted_categories_js
            ))
        return True
    
    def push_exports_to_github(self, export_dir="Web_Exports"):
        """将本地打包好的网页文件全自动推送到 GitHub 仓库"""
        token = os.getenv('GITHUB_TOKEN')
        repo_name = os.getenv('GITHUB_REPO') 

        if not token or not repo_name:
            return False, "发布失败：未在 .env 中找到 GitHub Token 或 仓库名"

        try:
            g = Github(token)
            repo = g.get_repo(repo_name)
            
            # 遍历 Web_Exports 文件夹内的所有生成文件
            for root, dirs, files in os.walk(export_dir):
                for file in files:
                    local_path = os.path.join(root, file)
                    # 将 Windows 本地路径转换为 GitHub 支持的斜杠相对路径
                    git_path = os.path.relpath(local_path, export_dir).replace('\\', '/')
                    
                    # 读取文件内容
                    with open(local_path, 'rb') as f:
                        content = f.read()
                        
                    try:
                        # 检查仓库中是否已有该文件，如果有则覆盖更新 (需提供 sha)
                        file_in_repo = repo.get_contents(git_path)
                        repo.update_file(file_in_repo.path, f"🤖 自动更新网页: {git_path}", content, file_in_repo.sha)
                    except:
                        # 如果仓库中没有该文件，则直接新建
                        repo.create_file(git_path, f"🚀 自动新增网页: {git_path}", content)
                        
            return True, "🎉 网页打包完成，已全自动推送到 GitHub 云端！"
            
        except Exception as e:
            return False, f"推送到 GitHub 时发生异常: {str(e)}"
